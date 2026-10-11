"""Generate both PDF lengths and a PPTX/PDF deck from recorded metrics."""
from pathlib import Path
import csv
import json
import sys
from html import escape
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from src.data_loader import ROOT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

DOCS=ROOT/'docs'
DOCS.mkdir(exist_ok=True)
S=json.loads((ROOT/'results/summary.json').read_text())
W=json.loads((ROOT/'results/weights.json').read_text())
TEAM=json.loads((ROOT/'team.json').read_text())
R=list(csv.DictReader((ROOT/'results/metrics.csv').open()))
A=S['selected_test'];B=S['baseline_test'];P=S['selected_by_validation']
NAMES=' | '.join(f"{m['name']} ({m['srn']})" for m in TEAM['members'])
METADATA=f"UE24CS352A Machine Learning | Section: {TEAM['section']} | Problem: {TEAM['problem_number']}"
def pct(n):return f'{100*n:.2f}%'

TEXT={
'Problem statement': 'Assign noisy vehicle GPS observations to directed road arcs in real time, without shortest-path searches. The objective is to implement the naive HMM approach of Ahres et al. and measure how delaying irrevocable predictions by N observations changes accuracy for a window of T observations.',
'Dataset': f"The public Newson-Krumm Seattle benchmark contains {S['dataset']['gps_points']:,} GPS readings, {S['dataset']['road_edges']:,} road edges ({S['dataset']['directed_arcs']:,} directed arcs), and a 576-edge reference route. The nominal 1 Hz trace has four sampling gaps and spans 7,611 seconds. Longitude/latitude are projected to UTM zone 10N. Speed and heading use a causal five-reading displacement. The first {S['training_points']:,} points train weights, the next {S['validation_points']:,} select T/N, and the final {S['test_points']:,} form the test set (60/10/30 split).",
'Reference-label construction': 'The benchmark supplies an ordered route, not independent per-point labels. Offline dynamic programming minimizes distance to route arcs under monotone progress, allowing stay/advance by up to eight edges and adding a one-metre penalty per advanced edge. Repeated route occurrences remain distinct. These derived labels are evaluation proxies; the median GPS-to-aligned-arc distance is 2.75 m. Candidate coverage on test is 99.91%.',
'Approach': 'Candidate states are all directed arcs within 50 m of a GPS point, found using a spatial index. Emissions combine Gaussian point-to-road density (sigma=10 m), directed heading agreement, and a high-speed/speed-limit constraint. Transitions combine a committed-history backtrack penalty and reciprocal GPS-versus-arc-midpoint distance discrepancy. Finite-window Viterbi maximizes summed log potentials, clamps already committed states at their actual indices, and commits state t-N when reading t arrives. At EOF, the remaining N states are flushed. No road-network shortest path or connectivity check is used.',
'Learning and implementation': f"Python modules separate loading, candidates, features, decoding, learning, evaluation, and visualization. Balanced L1 LinearSVC models learn emission then transition coefficients independently from reference labels. Ten expanding time folds, with a 15-observation gap, select C from 0.1, 1, 10 using F1 (selected emission C={W['emission_training']['C']:g}, transition C={W['transition_training']['C']:g}). Standardization is fitted within each fold. Positive raw-space coefficients are L1-normalized; negatives are clipped, the intercept is discarded, and scores have an epsilon floor. These are uncalibrated positive potentials. Training uses reference history for backtrack features; inference uses committed history.",
'Results': f"Validation selects T={P['T']}, N={P['N']} (validation directed error {pct(P['validation_arc_error'])}). On 2,260 held-out points, directed-arc error is {pct(A['arc_error'])}, edge error {pct(A['edge_error'])}, and online-only directed error {pct(A['online_arc_error'])}; seven EOF points are flushed. Nearest-road edge error is {pct(B['edge_error'])}. Directed baseline error is {pct(B['arc_error'])}, but its forward-direction tie rule disadvantages that baseline. Matching takes {A['seconds_per_1000']:.3f} s/1,000 points on this machine, excluding loading and visualization.",
'Conclusions': 'Buffering reduces error relative to zero delay, but the test minima differ by window: N=3, 8, and 8 for T=5, 10, and 15. The curves are not uniformly convex, and the minimum is not always T/2. Removing distance discrepancy improves directed error from 9.60% to 8.94%; this feature is not universally helpful. Other ablations worsen error. Results differ from the paper because the benchmark, labels, learned weights, and decoder choices differ. Disconnected predictions and proxy-label bias remain limitations. Future work should use independent point labels, additional routes, causal noise filtering, and better turn/connectivity features.'}

REFERENCES='Ahres, Jambulapati, Janssen & Kangaspunta, Real-Time Dense Map Matching with Naive Hidden Markov Models: Delay versus Accuracy, supplied CS229 paper. Newson & Krumm (2009), Hidden Markov Map Matching Through Noise and Sparseness, ACM GIS, pp. 336-343, doi:10.1145/1653771.1653818. Forney (1973), The Viterbi Algorithm, Proceedings of the IEEE 61(3), pp. 268-278.'

def styles(size=9.5):
    s=getSampleStyleSheet()
    s.add(ParagraphStyle(name='BodyCustom',fontName='Helvetica',fontSize=size,leading=size*1.27,spaceAfter=7))
    s.add(ParagraphStyle(name='SectionCustom',fontName='Helvetica-Bold',fontSize=size+1,leading=size+3,textColor=colors.HexColor('#153e75'),spaceBefore=6,spaceAfter=3))
    return s

def report_table(style):
    rows=[['Model','T / N','Directed error','Edge error']]
    for name,t,n in [('Zero delay',10,0),('Validation-selected',P['T'],P['N'])]:
        item=next(x for x in R if int(x['T'])==t and int(x['N'])==n)
        rows.append([name,f'{t} / {n}',pct(float(item['arc_error'])),pct(float(item['edge_error']))])
    rows.append(['Nearest segment','- / 0',pct(B['arc_error']),pct(B['edge_error'])])
    table=Table(rows,colWidths=[155,70,110,110])
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#153e75')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),8.5),('BOTTOMPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#edf2f7'),colors.white])]))
    return table

def header(s):
    return [Paragraph('Real-Time Dense Map Matching<br/>with Naive Hidden Markov Models',s['Title']),
            Paragraph(escape(NAMES),s['BodyCustom']),
            Paragraph(escape(METADATA),s['BodyCustom'])]

def section(story,s,title,body):
    story.extend([Paragraph(escape(title),s['SectionCustom']),Paragraph(escape(body),s['BodyCustom'])])

def footer(canvas,doc):
    canvas.setFont('Helvetica',8);canvas.setFillColor(colors.grey)
    canvas.drawString(36,22,'UE24CS352A | Measured Seattle benchmark results | Prepared 12 October 2026')
    canvas.drawRightString(A4[0]-36,22,str(doc.page))

def reports():
    s=styles(9.5);story=header(s)
    for key in ['Problem statement','Dataset','Reference-label construction','Approach']:
        section(story,s,key,TEXT[key])
    story.extend([Image(str(ROOT/'results/error_vs_buffer.png'),width=460,height=223),Spacer(1,5),report_table(s),PageBreak()])
    for key in ['Learning and implementation','Results']:
        section(story,s,key,TEXT[key])
    story.extend([Image(str(ROOT/'results/ablation.png'),width=460,height=197)])
    section(story,s,'Conclusions',TEXT['Conclusions'])
    section(story,s,'Reproducibility','Install requirements; download the benchmark; run python -m src.train_weights, python -m src.evaluate, and python run_demo.py --T 10 --N 7. Saved CSV metrics feed this PDF and the deck. Source repository URL has not been supplied; the GitHub deliverable still requires publishing and collaborator access.')
    section(story,s,'References',REFERENCES)
    SimpleDocTemplate(str(DOCS/'writeup.pdf'),pagesize=A4,leftMargin=36,rightMargin=36,topMargin=30,bottomMargin=36).build(story,onFirstPage=footer,onLaterPages=footer)
    s=styles(8.5);story=header(s)
    section(story,s,'Problem',TEXT['Problem statement'])
    section(story,s,'Dataset',TEXT['Dataset']+' Per-point truth is an offline monotone route-distance alignment, not independent ground truth; test candidate coverage is 99.91%.')
    section(story,s,'Approach and implementation',TEXT['Approach']+' Balanced L1-SVM weights are fitted using ten expanding training-only time folds (15-point gap); positive coefficients become normalized potentials and intercepts are discarded. Python, Shapely, pyproj, NumPy and scikit-learn implement the pipeline. Run python run_demo.py --T 10 --N 7 after training.')
    story.extend([Image(str(ROOT/'results/error_vs_buffer.png'),width=420,height=204),report_table(s)])
    section(story,s,'Results and conclusions',f"The validation-selected model runs in {A['seconds_per_1000']:.3f} s/1,000 points, excluding loading. Seven EOF predictions are flushed. Buffers improve accuracy, but minima are not consistently at T/2 and curves are not uniformly convex. Dropping distance discrepancy improves directed error to 8.94%. Proxy labels, one trip, sampling gaps and disconnected predictions limit conclusions. Future work: independent labels, more routes, turn features and causal filtering. GitHub URL is not supplied; publishing remains pending.")
    section(story,s,'References', 'Ahres et al., supplied CS229 paper; Newson & Krumm (2009), ACM GIS, doi:10.1145/1653771.1653818; Forney (1973), The Viterbi Algorithm.')
    SimpleDocTemplate(str(DOCS/'writeup_one_page.pdf'),pagesize=A4,leftMargin=36,rightMargin=36,topMargin=24,bottomMargin=32).build(story,onFirstPage=footer,onLaterPages=footer)
    md='# Real-Time Dense Map Matching with Naive Hidden Markov Models\n\n'+NAMES+'\n\n'
    md+='\n\n'.join('## '+key+'\n\n'+value for key,value in TEXT.items())+'\n\n## References\n\n'+REFERENCES+'\n'
    (DOCS/'writeup_source.md').write_text(md,encoding='utf-8')

SLIDES=[
('Real-Time Dense Map Matching', ['Naive hidden Markov models: delay versus accuracy', *[m['name']+' | '+m['srn'] for m in TEAM['members']], 'UE24CS352A Machine Learning'],None),
('Problem and motivation',['GPS noise can place a vehicle on the wrong road.','Estimate the hidden directed road arc for each reading.','Waiting N readings provides future evidence but delays output.','Avoid shortest-path searches; study the measured trade-off.'],None),
('Seattle benchmark',['7,531 readings; 158,167 road edges; 576 reference-route edges.','Nominal 1 Hz with four gaps; coordinates projected to metres.','60% train / 10% validation / 30% held-out test.','Point labels derived from route alignment: evaluation proxies.'],'map.png'),
('Pipeline',['GPS + roads -> nearby directed arc candidates','Distance, heading, speed features -> emissions','History and displacement features -> transitions','Log-domain sliding-window Viterbi -> irrevocable predictions'],None),
('HMM and features',['States: all directed arcs within 50 m; sigma = 10 m.','Emission: Gaussian distance, heading dot product, speed constraint.','Transition: backtrack prevention and midpoint discrepancy.','Positive weighted potentials; no connectivity guarantee.'],None),
('Window T and buffer N',['At reading t, solve over max(0, t-T+1) through t.','Clamp overlapping committed states at their actual times.','Commit prediction for t-N; never revise it later.','Flush final N points at EOF; report online-only error separately.'],None),
('Learning feature weights',['Balanced L1 LinearSVC; candidates/pairs labelled from reference.','10 expanding folds, 15-reading gap; C in {0.1, 1, 10}.',f"Selected C: emission {W['emission_training']['C']:g}, transition {W['transition_training']['C']:g}.",'Positive coefficients normalized; intercept discarded.'],None),
('Delay versus accuracy',[f"Validation selected T={P['T']}, N={P['N']}.",f"Test directed error: {pct(A['arc_error'])}.",f"Test edge error: {pct(A['edge_error'])}.",'Minimum is not always at T/2; curves are not uniformly convex.'],'error_vs_buffer.png'),
('Baseline and ablations',[f"Nearest-road edge error: {pct(B['edge_error'])}.",f"Full model edge error: {pct(A['edge_error'])}.",'Removing heading or distance sharply worsens error.','Removing discrepancy improves directed error to 8.94%.'],'ablation.png'),
('Live demonstration',['python run_demo.py --T 10 --N 7 --slice 5300 5600','Open results/demo_map.html: completely offline replay.','Compare N=0 and N=7 using the same slice.','Static PNG is the backup; runtime excludes map rendering.'],'demo_map.png'),
('Conclusions and limitations',['Buffering and heading help on this single benchmark.','Derived labels and one trip limit generalization.','Training/inference history differ; no network connectivity check.','Next: independent labels, multiple trips, filtering, turn features.'],None),
('Code, contribution and Q&A',['src/: loader, candidates, features, Viterbi, training, evaluation.','Tests verify causal commitments and future-evidence effects.','Both members should review the full pipeline and run the demo.','Record actual individual contributions; private GitHub publishing pending.'],None)]

def deck():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    pdf=Canvas(str(DOCS/'presentation.pdf'),pagesize=(960,540))
    for number,(title,bullets,img) in enumerate(SLIDES,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(15,23,42)
        box=slide.shapes.add_textbox(Inches(.6),Inches(.5),Inches(12),Inches(.9))
        p=box.text_frame.paragraphs[0];p.text=title;p.font.size=Pt(30);p.font.bold=True;p.font.color.rgb=RGBColor(248,250,252)
        width=5 if img else 11.8
        box=slide.shapes.add_textbox(Inches(.65),Inches(1.75),Inches(width),Inches(4.7))
        tf=box.text_frame;tf.word_wrap=True
        for i,bullet in enumerate(bullets):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=bullet;p.font.size=Pt(20 if img else 24);p.font.color.rgb=RGBColor(203,213,225);p.space_after=Pt(22)
        if img:
            path=ROOT/'results'/img
            from PIL import Image as PILImage
            iw,ih=PILImage.open(path).size
            image_width=6.6;image_height=image_width*ih/iw
            slide.shapes.add_picture(str(path),Inches(6),Inches(1.8),width=Inches(image_width),height=Inches(image_height))
        box=slide.shapes.add_textbox(Inches(.6),Inches(7),Inches(12),Inches(.3))
        p=box.text_frame.paragraphs[0];p.text=f'UE24CS352A  |  Harish & Sharanabasava  |  {number:02d} / 12';p.font.size=Pt(11);p.font.color.rgb=RGBColor(148,163,184)
        slide.notes_slide.notes_text_frame.text='\n'.join(bullets)+'\nReference: supplied Ahres et al. CS229 paper; Newson and Krumm Seattle benchmark. Explain implementation adaptations honestly. Contribution assignments are not verified.'
        pdf.setFillColor(colors.HexColor('#0f172a'));pdf.rect(0,0,960,540,fill=1,stroke=0)
        pdf.setFillColor(colors.white);pdf.setFont('Helvetica-Bold',25);pdf.drawString(44,477,title)
        ss=ParagraphStyle('slide',fontName='Helvetica',fontSize=15 if img else 19,leading=23 if img else 28,textColor=colors.HexColor('#cbd5e1'))
        y=415
        for bullet in bullets:
            para=Paragraph(escape(bullet),ss);pw,ph=para.wrap(360 if img else 845,400);para.drawOn(pdf,48,y-ph);y-=ph+22
        if img:
            pdf.drawImage(str(ROOT/'results'/img),432,120,width=475,height=290,preserveAspectRatio=True,anchor='c',mask='auto')
        pdf.setFont('Helvetica',9);pdf.setFillColor(colors.HexColor('#94a3b8'));pdf.drawString(44,22,f'UE24CS352A | Harish & Sharanabasava | {number:02d} / 12');pdf.showPage()
    prs.save(DOCS/'presentation.pptx');pdf.save()

def qa():
    body=f'''# Live demo and Q&A

## Rehearsal sequence (about three minutes)

1. Run `python run_demo.py --T 10 --N 0 --slice 5300 5600`; note printed errors and open `results/demo_map.html`.
2. Run `python run_demo.py --T 10 --N 7 --slice 5300 5600`; compare identical points. Reload the HTML (the file is overwritten).
3. Show `results/error_vs_buffer.png` and explain that T=10/N=7 was selected on validation, not test.
4. Show `results/ablation.png`, including the improvement after removing discrepancy.
5. Use `results/demo_map.png` or `docs/presentation.pdf` as offline backups. HTML requires no tiles or internet. Loading the first time builds a local cache and may take longer; pre-run once before review.

## Answers both members should understand

- **States and observations?** States are directed arcs; observations are projected GPS readings with causal speed/heading estimates.
- **Why an HMM?** Distance alone ignores sequence context; transitions encourage stable, plausible motion, and Viterbi searches the sequence of candidates.
- **What does N mean?** Commit point t-N when point t arrives; N is future readings, about N seconds except across four gaps.
- **Does N always help?** No. Our held-out curves have minima/plateaus, not uniformly convex curves. More delay reduces available committed past inside a fixed window.
- **Why no shortest path?** This project tests the paper's inexpensive local feature approximation. Nearby arcs are not guaranteed reachable; disconnected predictions are a known failure.
- **Exact feature equations?** E1 = exp(-d²/(2 sigma²))/(sigma sqrt(2 pi)); E2 = heading dot local directed tangent; E3 = 0 above 40 mph and >=1.2 road speed, else 1. T1 = 0 on already committed arcs except staying on current arc; T2 = 1/(abs(GPS displacement - midpoint separation)+1).
- **Are these probabilities?** No. SVM coefficients become positive potentials after clipping/normalization; we omit the classifier intercept and use an epsilon floor. This is a disclosed adaptation.
- **What does L1 do?** Encourages sparse coefficients. Balanced class weights correct many more negative candidates than positives. Standardization is fitted on each training fold only.
- **How is leakage prevented?** First 60% training, next 10% selection, final 30% test; ten expanding training folds with a 15-reading gap. GPS features are causal. Route labels are offline references, never decoder inputs. One trip and related roads can still bias evaluation.
- **Where did truth come from?** The published ordered route is aligned to points using monotone distance DP. These are derived proxy labels, not independent surveyed point truth. Inspect both map and route-progress diagnostics.
- **What are the measured results?** T={P['T']}, N={P['N']}: directed error {pct(A['arc_error'])}, edge error {pct(A['edge_error'])}; nearest-road edge error {pct(B['edge_error'])}. Matcher {A['seconds_per_1000']:.3f} s/1,000 points on this machine; loading/plotting excluded.
- **Why report edge error?** It ignores direction; a geometric nearest-road method cannot resolve equal-distance forward/reverse arcs. Its directed error depends on an arbitrary forward tie rule.
- **Why does dropping discrepancy help?** Whole-road midpoint distance can be a poor displacement proxy for long/curved edges. The test result is descriptive, not used to retroactively select the submitted model.
- **How is backtrack history handled?** A visited set of committed states in the current window preserves first-order Viterbi. Training uses route-label history; inference uses its own commitments. This mismatch is disclosed.
- **Why does T sometimes matter little?** Overlapping committed states are fixed; only N+1 states remain revisable. T also controls the available history for backtrack penalties.
- **What is EOF flush?** Final N states cannot receive full future context. Full error includes them; online-only error excludes them.
- **Who did what?** Use the actual record in docs/contributions.md. Do not claim fictional authorship or artificial commit history.

## Submission actions

Publish all source to a private GitHub repository; add the faculty and assigned TAs, and verify access. Provide section, assigned problem number and repository URL in team.json, then regenerate documents. Submit the PDF; use the PPTX/PDF deck and demonstrate live. The assignment's stated deadline is October 10, 2026, 11:59 PM; confirm current acceptance arrangements with faculty.
'''
    (DOCS/'demo_and_qa.md').write_text(body,encoding='utf-8')

if __name__=='__main__':
    reports();deck();qa()
    print('Generated two-page and one-page write-ups, 12-slide PPTX/PDF, source text and Q&A.')
